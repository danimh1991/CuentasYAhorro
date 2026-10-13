"""Export the Bolsa movement ledger added after the initial workbook import."""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import openpyxl


SOURCE = Path(os.environ["TASK_XLSX"])
OUTPUT = Path(os.environ.get("TASK_SQL_OUTPUT", ".local/excel_phase2.sql"))
WB = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def cents(value) -> int:
    return round(value * 100) if isinstance(value, (int, float)) else 0


name_map = {"Nasqad 100": "Nasdaq 100"}
sheet = WB["Bolsa"]
lines = ["PRAGMA foreign_keys = ON;", "BEGIN TRANSACTION;"]

for row in range(56, sheet.max_row + 1):
    trade_date = sheet.cell(row, 16).value
    raw_name = sheet.cell(row, 17).value
    units = sheet.cell(row, 18).value
    price = sheet.cell(row, 19).value
    total = sheet.cell(row, 20).value
    kind = sheet.cell(row, 25).value
    broker = sheet.cell(row, 26).value
    if not isinstance(trade_date, (datetime, date)) or not raw_name or not isinstance(units, (int, float)):
        continue
    if kind not in {"A", "D", "R", "E"}:
        continue
    name = name_map.get(str(raw_name), str(raw_name))
    instrument_lookup = (
        "(SELECT id FROM instruments WHERE category='stocks' "
        "AND name IN ('Nasdaq 100', 'Nasqad 100') ORDER BY name LIMIT 1)"
        if name == "Nasdaq 100"
        else f"(SELECT id FROM instruments WHERE category='stocks' AND name={sql(name)})"
    )
    key = f"excel-bolsa-{row}"
    lines.append(
        "INSERT OR IGNORE INTO stock_transactions "
        "(source_key, trade_date, instrument_id, kind, units, price_minor, price_scale, total_cents, currency, fx_to_eur, broker, source) "
        f"VALUES ({sql(key)}, {sql(trade_date.strftime('%Y-%m-%d'))}, "
        f"{instrument_lookup}, {sql(kind)}, {units}, "
        f"{round((price or 0) * 10000)}, 10000, {cents(total)}, 'EUR', 1, {sql(broker)}, 'Excel');"
    )

lines.append("COMMIT;")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Generated {OUTPUT} with {len(lines) - 3} stock movements")
