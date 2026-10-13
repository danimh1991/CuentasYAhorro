from pathlib import Path
import openpyxl

SOURCE = Path(r"\\Danieta_Nas\home\Drive\Documentos Personales\Cuentas_Ahorro.xlsx")
OUTPUT = Path(".local/excel_phase4.sql")


def sql_text(value: object) -> str:
    return "'" + str(value).replace("'", "''") + "'"


book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)
sheet = book["Bolsa"]
lines = ["PRAGMA foreign_keys = ON;"]
count = 0

for row in range(56, sheet.max_row + 1):
    payment_date = sheet.cell(row, 28).value  # AB
    instrument = sheet.cell(row, 29).value  # AC
    per_share = sheet.cell(row, 30).value  # AD
    total = sheet.cell(row, 31).value  # AE
    if not payment_date or not instrument or not isinstance(total, (int, float)):
        continue
    date_text = payment_date.strftime("%Y-%m-%d")
    per_share_value = float(per_share) if isinstance(per_share, (int, float)) else 0
    lines.append(
        "INSERT OR REPLACE INTO stock_dividends "
        "(source_key, payment_date, instrument_id, gross_minor, amount_scale, currency, fx_to_eur, "
        "amount_eur_cents, note, source, per_share_minor, per_share_scale) VALUES "
        f"({sql_text(f'excel-dividend-{row}')}, {sql_text(date_text)}, "
        f"(SELECT id FROM instruments WHERE category='stocks' AND name={sql_text(instrument)}), "
        f"{round(float(total) * 10000)}, 10000, 'EUR', 1, {round(float(total) * 100)}, "
        f"'Histórico importado', 'Excel', {round(per_share_value * 1000000)}, 1000000);"
    )
    count += 1

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Generated {OUTPUT} with {count} dividends")
