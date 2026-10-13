"""Export Detalle Ahorro and Liberty summary metrics to a git-ignored D1 file."""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import openpyxl


SOURCE = Path(os.environ["TASK_XLSX"])
OUTPUT = Path(os.environ.get("TASK_SQL_OUTPUT", ".local/excel_phase3.sql"))
WB = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def cents(value) -> int:
    return round(value * 100) if isinstance(value, (int, float)) else 0


def day(value: datetime | date) -> str:
    return value.strftime("%Y-%m-%d")


lines = ["PRAGMA foreign_keys = ON;"]
counts = {"savings": 0, "fund_transactions": 0, "liberty": 0}

savings = WB["Detalle Ahorro"]
buckets = {
    "cash": (2, 4, None),
    "deposit": (7, 9, 10),
    "funds": (13, 15, None),
    "indexed": (18, 20, None),
}
for row in range(3, savings.max_row + 1):
    month = savings.cell(row, 1).value
    if not isinstance(month, (datetime, date)):
        continue
    month_key = month.strftime("%Y-%m")
    for bucket, (contribution_col, current_col, interest_col) in buckets.items():
        contribution = savings.cell(row, contribution_col).value
        current = savings.cell(row, current_col).value
        interest = savings.cell(row, interest_col).value if interest_col else 0
        if not any(isinstance(value, (int, float)) for value in (contribution, current, interest)):
            continue
        lines.append(
            "INSERT OR REPLACE INTO savings_entries "
            "(period_id, bucket, contribution_cents, current_value_cents, interest_cents, source) "
            f"VALUES ((SELECT id FROM periods WHERE substr(month,1,7)={sql(month_key)}), {sql(bucket)}, "
            f"{cents(contribution)}, {cents(current)}, {cents(interest)}, 'Excel');"
        )
        counts["savings"] += 1

for row in range(2, savings.max_row + 1):
    fund = savings.cell(row, 30).value
    trade_date = savings.cell(row, 31).value
    units = savings.cell(row, 33).value
    amount = savings.cell(row, 34).value
    price = savings.cell(row, 35).value
    if not fund or not isinstance(trade_date, (datetime, date)) or not isinstance(units, (int, float)):
        continue
    lines.append(
        "INSERT OR IGNORE INTO savings_fund_transactions "
        "(source_key, fund_name, trade_date, units, amount_cents, price_minor, price_scale, source) "
        f"VALUES ({sql(f'excel-ahorro-{row}')}, {sql(fund)}, {sql(day(trade_date))}, {units}, "
        f"{cents(amount)}, {round((price or 0) * 10000)}, 10000, 'Excel');"
    )
    counts["fund_transactions"] += 1

liberty = WB["Detalle Liberty"]
for row in range(1, liberty.max_row + 1):
    month = liberty.cell(row, 2).value
    if liberty.cell(row, 3).value != "Total:" or not isinstance(month, (datetime, date)):
        continue
    total = liberty.cell(row, 4).value
    gain = liberty.cell(row, 5).value
    monthly_gain = liberty.cell(row, 7).value
    return_rate = liberty.cell(row, 9).value
    contributed = (total - gain) if isinstance(total, (int, float)) and isinstance(gain, (int, float)) else 0
    lines.append(
        "INSERT OR REPLACE INTO liberty_snapshots "
        "(period_id, contributed_cents, gain_cents, monthly_gain_cents, return_rate, source) "
        f"VALUES ((SELECT id FROM periods WHERE substr(month,1,7)={sql(month.strftime('%Y-%m'))}), "
        f"{cents(contributed)}, {cents(gain)}, {cents(monthly_gain)}, "
        f"{return_rate if isinstance(return_rate, (int, float)) else 'NULL'}, 'Excel');"
    )
    counts["liberty"] += 1

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Generated {OUTPUT}: {counts}")
