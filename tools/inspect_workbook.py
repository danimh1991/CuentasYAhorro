import json
import os
import sys
from pathlib import Path

import openpyxl


SOURCE = Path(os.environ["TASK_XLSX"])
FORMULAS = openpyxl.load_workbook(SOURCE, data_only=False, read_only=False)
VALUES = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)
sys.stdout.reconfigure(encoding="utf-8")


def populated_rows(sheet_name: str, max_rows: int | None = None) -> list[dict]:
    formula_sheet = FORMULAS[sheet_name]
    value_sheet = VALUES[sheet_name]
    rows: list[dict] = []
    limit = min(formula_sheet.max_row, max_rows or formula_sheet.max_row)
    for row_idx in range(1, limit + 1):
        cells = []
        for col_idx in range(1, formula_sheet.max_column + 1):
            formula_cell = formula_sheet.cell(row_idx, col_idx)
            value_cell = value_sheet.cell(row_idx, col_idx)
            if formula_cell.value is None and value_cell.value is None:
                continue
            cells.append(
                {
                    "cell": formula_cell.coordinate,
                    "formula": formula_cell.value if formula_cell.data_type == "f" else None,
                    "value": value_cell.value,
                    "raw": None if formula_cell.data_type == "f" else formula_cell.value,
                    "format": formula_cell.number_format,
                }
            )
        if cells:
            rows.append({"row": row_idx, "cells": cells})
    return rows


if __name__ == "__main__":
    requested_range = os.environ.get("TASK_RANGE")
    if requested_range:
        sheet_name, coordinates = requested_range.split("!", 1)
        formula_sheet = FORMULAS[sheet_name]
        value_sheet = VALUES[sheet_name]
        output = []
        for formula_row, value_row in zip(formula_sheet[coordinates], value_sheet[coordinates]):
            output.append([
                {
                    "cell": formula_cell.coordinate,
                    "value": value_cell.value,
                    "formula": formula_cell.value if formula_cell.data_type == "f" else None,
                }
                for formula_cell, value_cell in zip(formula_row, value_row)
            ])
        print(json.dumps(output, ensure_ascii=False, indent=2, default=str))
        raise SystemExit
    selected = os.environ.get("TASK_SHEETS", "").split("|")
    max_rows = int(os.environ["TASK_MAX_ROWS"]) if os.environ.get("TASK_MAX_ROWS") else None
    result = {name: populated_rows(name, max_rows) for name in selected if name}
    tail = int(os.environ.get("TASK_TAIL", "0"))
    if tail:
        result = {name: rows[-tail:] for name, rows in result.items()}
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
