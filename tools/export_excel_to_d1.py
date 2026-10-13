"""Convert the private Excel workbook into a local, git-ignored D1 import file."""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import openpyxl


SOURCE = Path(os.environ["TASK_XLSX"])
OUTPUT = Path(os.environ.get("TASK_SQL_OUTPUT", ".local/excel_import.sql"))
WB = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def iso_day(value: datetime | date) -> str:
    return value.strftime("%Y-%m-%d")


def cents(value) -> int | None:
    return None if not isinstance(value, (int, float)) else round(value * 100)


lines = ["PRAGMA foreign_keys = ON;"]
use_transaction = os.environ.get("TASK_NO_TRANSACTION") != "1"
if use_transaction:
    lines.append("BEGIN TRANSACTION;")

valor = WB["Valor"]
period_dates: list[str] = []
for row in range(2, valor.max_row + 1):
    month = valor.cell(row, 1).value
    if not isinstance(month, (datetime, date)):
        continue
    month_text = iso_day(month)
    period_dates.append(month_text)
    values = [cents(valor.cell(row, column).value) for column in range(2, 9)]
    current, savings, pias_a, pias_b, liberty, pension, stocks = values
    pias = (pias_a or 0) + (pias_b or 0)
    total = sum(value or 0 for value in (current, savings, pias, liberty, pension, stocks))
    note = valor.cell(row, 12).value
    lines.append(
        f"INSERT OR IGNORE INTO periods (month, status, notes, closed_at) VALUES ({sql(month_text)}, 'closed', {sql(note)}, CURRENT_TIMESTAMP);"
    )
    lines.append(
        "INSERT OR REPLACE INTO summary_snapshots "
        "(period_id, current_account_cents, savings_account_cents, pias_cents, liberty_cents, pension_cents, stocks_cents, total_cents) "
        f"VALUES ((SELECT id FROM periods WHERE month={sql(month_text)}), {sql(current)}, {sql(savings)}, {sql(pias)}, "
        f"{sql(liberty)}, {sql(pension)}, {sql(stocks)}, {sql(total)});"
    )

liberty_sheet = WB["Detalle Liberty"]
for total_row in range(1, liberty_sheet.max_row + 1):
    month = liberty_sheet.cell(total_row, 2).value
    total_value = liberty_sheet.cell(total_row, 4).value
    label = liberty_sheet.cell(total_row, 3).value
    if label != "Total:" or not isinstance(month, (datetime, date)) or not isinstance(total_value, (int, float)):
        continue
    month_text = iso_day(month)
    header_row = total_row - 1
    while header_row > 0 and liberty_sheet.cell(header_row, 1).value != "Fondos":
        header_row -= 1
    for row in range(header_row + 1, total_row):
        name = liberty_sheet.cell(row, 1).value
        units = liberty_sheet.cell(row, 2).value
        market_value = liberty_sheet.cell(row, 4).value
        price = liberty_sheet.cell(row, 6).value
        if not name or not isinstance(units, (int, float)) or not isinstance(market_value, (int, float)):
            continue
        lines.append(
            "INSERT OR IGNORE INTO instruments (category, name, sort_order) "
            f"VALUES ('liberty', {sql(name)}, {row - header_row});"
        )
        price_minor = round((price or (market_value / units if units else 0)) * 10000)
        lines.append(
            "INSERT OR REPLACE INTO holdings "
            "(period_id, instrument_id, units, price_minor, price_scale, price_currency, fx_to_eur, price_eur_minor, market_value_cents, quote_source, quoted_at) "
            f"VALUES ((SELECT id FROM periods WHERE month={sql(month_text)}), (SELECT id FROM instruments WHERE category='liberty' AND name={sql(name)}), "
            f"{units}, {price_minor}, 10000, 'EUR', 1, {price_minor}, {cents(market_value)}, 'Importado de Excel', {sql(month_text + 'T12:00:00Z')});"
        )

pension_sheet = WB["Fondo Pensiones HPE"]
for row in range(3, pension_sheet.max_row + 1):
    month = pension_sheet.cell(row, 1).value
    if not isinstance(month, (datetime, date)):
        continue
    month_text = iso_day(month)
    employer = cents(pension_sheet.cell(row, 2).value) or 0
    personal = cents(pension_sheet.cell(row, 3).value) or 0
    current = cents(pension_sheet.cell(row, 4).value) or 0
    cumulative = cents(pension_sheet.cell(row, 5).value) or 0
    lines.append(
        "INSERT OR REPLACE INTO pension_entries "
        f"(period_id, employer_cents, personal_cents, current_value_cents, cumulative_contributions_cents) VALUES "
        f"((SELECT id FROM periods WHERE month={sql(month_text)}), {employer}, {personal}, {current}, {cumulative});"
    )

if period_dates:
    latest_month = period_dates[-1]
    bolsa = WB["Bolsa"]
    for row in range(3, 18):
        name = bolsa.cell(row, 1).value
        market_value = bolsa.cell(row, 2).value
        units = bolsa.cell(row, 3).value
        invested = bolsa.cell(row, 4).value
        withdrawn = bolsa.cell(row, 6).value
        if not name or not isinstance(units, (int, float)):
            continue
        lines.append(
            "INSERT OR IGNORE INTO instruments (category, name, sort_order) "
            f"VALUES ('stocks', {sql(name)}, {(row - 2) * 10});"
        )
        price = market_value / units if units and isinstance(market_value, (int, float)) else 0
        lines.append(
            "INSERT OR REPLACE INTO holdings "
            "(period_id, instrument_id, units, invested_cents, withdrawn_cents, price_minor, price_scale, price_currency, fx_to_eur, price_eur_minor, market_value_cents, quote_source, quoted_at) "
            f"VALUES ((SELECT id FROM periods WHERE month={sql(latest_month)}), (SELECT id FROM instruments WHERE category='stocks' AND name={sql(name)}), "
            f"{units}, {sql(cents(invested))}, {sql(cents(withdrawn) or 0)}, {round(price * 10000)}, 10000, 'EUR', 1, {round(price * 10000)}, "
            f"{sql(cents(market_value))}, 'Importado de Excel', {sql(latest_month + 'T12:00:00Z')});"
        )

if use_transaction:
    lines.append("COMMIT;")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Generated {OUTPUT} with {len(lines)} statements")
